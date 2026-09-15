"""TRACE intake channel: **PORB** (Plan of Results and Budget, 2025 master workbook).

Turns the approved public PORB master spreadsheet into the results-framework objects of a
Science Program: areas of work, high-level outputs, KPI/indicator lines, ToC outcomes,
W3/bilateral projects, MELIA studies, partners, and the geography of implementation and
benefit.

    python -m ingest.porb --xlsx <master.xlsx> --program "Breeding for Tomorrow" \
        --out data/claims/seed-porb.jsonl --extract data/seed/porb_sp01.json

Money boundary (CONTRACT §3.6 / Jose's concern (c)): the **HLO budget is an attribute on the
HLO object**, computed as the sum of its KPI rows. It is never expressed as a link and never
propagated to neighbours. Partner budgets and W3/bilateral amounts stay on the row-level
object/link they were reported on (``WITH_PARTNER.attrs.budget_usd``,
``FUNDED_BY.attrs.amount_usd``) and are explicitly *not* summable with HLO budgets.

Assumptions:

* The Partners sheet has no HLO column — partners are reported at **AoW** granularity, so we
  emit ``WITH_PARTNER aow -> institution`` (``attrs.granularity="aow"``) rather than inventing
  an HLO attribution.
* The W3-Bilateral sheet's "High Level Output title" column in practice contains **outcome**
  codes (``I-OC 3.5 …``) as often as HLO codes; we link ``FUNDED_BY`` from whichever object the
  text resolves to and record ``attrs.matched_on``.
* Rows whose label contains "subtotal"/"total" are skipped (they are spreadsheet artefacts).
* Geography values that are regions ("Western Africa") become ``region`` objects; "Global"
  is recorded as ``attrs.global_share_pct`` on the AoW instead of a link.
"""
from __future__ import annotations

import argparse
import re
from collections import defaultdict
from pathlib import Path
from typing import Any, Iterable

from ingest import common as C

CHANNEL = "porb"
SYSTEM = "porb"

DEFAULT_XLSX = ("/Users/smithai/workspace/PORBs_consolidation/SummaryPORBs_v3/"
                "PORB_MASTER_All_Programs_13Aug2026.xlsx")

#: PORB program/accelerator names -> CGIAR portfolio codes (CLARISA initiative codes 50-62).
PROGRAM_CODES: dict[str, str] = {
    "breeding for tomorrow": "SP01",
    "sustainable farming": "SP02",
    "sustainable animal and aquatic foods": "SP03",
    "sustainable aquatic and animal foods": "SP03",
    "multifunctional landscapes": "SP04",
    "better diets and nutrition": "SP05",
    "climate action": "SP06",
    "policy innovations": "SP07",
    "food frontiers and security": "SP08",
    "scaling for impact": "SP09",
    "gender equality and inclusion": "SP10",
    "capacity sharing": "SP11",
    "digital transformation": "SP12",
    "genebank": "SP13",
    "genebanks": "SP13",
}

UN_REGIONS = {
    "eastern africa", "western africa", "middle africa", "southern africa", "northern africa",
    "sub-saharan africa", "africa", "eastern asia", "southern asia", "south-eastern asia",
    "western asia", "central asia", "asia", "latin america and the caribbean",
    "central america", "south america", "northern america", "caribbean", "europe", "oceania",
    "melanesia", "polynesia", "micronesia", "west and central africa", "east and southern africa",
}

SUBTOTAL_RE = re.compile(r"\b(sub)?total\b", re.I)
AOW_RE = re.compile(r"^\s*(AOW\d+)\s*[:\-–]\s*(.+?)\s*$", re.I)
HLO_RE = re.compile(r"^\s*(HLO\s*\d+(?:\.\w+)*)\s+(.*)$", re.I)
OUTCOME_RE = re.compile(r"^\s*((?:I-OC|2030-OC|EoI-OC)\s*[\d.]+)\s*\.?\s*(.*)$", re.I | re.S)


def clean(value: Any) -> str | None:
    if value is None:
        return None
    s = re.sub(r"\s+", " ", str(value).replace("\xa0", " ")).strip()
    return s or None


def money(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return round(float(value), 2)
    except (TypeError, ValueError):
        return None


def is_subtotal(*cells: Any) -> bool:
    return any(SUBTOTAL_RE.search(str(c)) for c in cells if c)


def split_list(value: Any) -> list[str]:
    if not value:
        return []
    return [p.strip() for p in re.split(r"[;,]", str(value).replace("\xa0", " ")) if p.strip()]


# --------------------------------------------------------------------------------------
def read_sheets(xlsx: str, program: str) -> dict[str, list[dict[str, Any]]]:
    """Read every sheet, keep the rows of one program, return dicts keyed by header."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    out: dict[str, list[dict[str, Any]]] = {}
    programs: set[str] = set()
    try:
        for name in wb.sheetnames:
            ws = wb[name]
            rows = ws.iter_rows(values_only=True)
            header = [clean(h) or f"col{i}" for i, h in enumerate(next(rows, []) or [])]
            kept = []
            for raw in rows:
                if not raw or raw[0] is None:
                    continue
                prog = clean(raw[0])
                if prog:
                    programs.add(prog)
                if prog != program:
                    continue
                kept.append({header[i]: raw[i] for i in range(min(len(header), len(raw)))})
            out[name] = kept
        out["_programs"] = [{"name": p} for p in sorted(programs)]
    finally:
        wb.close()
    return out


# --------------------------------------------------------------------------------------
def parse_aow(value: Any, program_code: str) -> dict[str, Any] | None:
    text = clean(value)
    if not text:
        return None
    m = AOW_RE.match(text)
    code, label = (m.group(1).upper(), m.group(2)) if m else (text.split(":")[0].upper(), text)
    return {"id": C.make_id("aow", f"{program_code.lower()}-{code.lower()}"),
            "code": code, "label": label, "raw": text}


def parse_hlo(value: Any, program_code: str, aow_code: str | None) -> dict[str, Any] | None:
    text = clean(value)
    if not text or is_subtotal(text):
        return None
    m = HLO_RE.match(text)
    if m:
        code = re.sub(r"\s+", "", m.group(1)).upper()
        title = m.group(2).strip() or code
    else:
        code, title = None, text
    key = (C.slugify(code) if code else f"x-{C.sha1_12(text)}")
    return {"id": C.make_id("hlo", f"{program_code.lower()}-{key}"),
            "code": code, "label": title, "raw": text, "aow_code": aow_code}


def parse_outcome(value: Any, program_code: str) -> dict[str, Any] | None:
    text = clean(value)
    if not text or is_subtotal(text):
        return None
    m = OUTCOME_RE.match(text)
    if m:
        code = re.sub(r"\s+", " ", m.group(1)).strip().upper()
        title = m.group(2).strip(" .") or code
        key = C.slugify(code)
    else:
        code, title, key = None, text, f"x-{C.sha1_12(text)}"
    return {"id": C.make_id("outcome", f"{program_code.lower()}-{key}"),
            "code": code, "label": title[:400], "raw": text}


# --------------------------------------------------------------------------------------
def run(xlsx: str, program: str, out: str, extract: str,
        prms_extract: str | None = None) -> dict[str, Any]:
    snapshot = Path(xlsx).stem
    program_code = PROGRAM_CODES.get(program.lower().strip(), "SP01")
    sheets = read_sheets(xlsx, program)
    matcher = C.InstitutionMatcher(C.load_prms_institutions(prms_extract or ""))

    claims: list[dict[str, Any]] = []
    seen: set[str] = set()

    def obj(*args, ref: str, **kwargs) -> str:
        claim = C.object_claim(*args, channel=CHANNEL, system=SYSTEM, ref=ref,
                               snapshot=snapshot, **kwargs)
        oid = claim["payload"]["id"]
        if oid not in seen:
            seen.add(oid)
            claims.append(claim)
        return oid

    def link(subject: str, predicate: str, target: str, ref: str, **kwargs) -> None:
        claims.append(C.link_claim(subject, predicate, target, channel=CHANNEL, system=SYSTEM,
                                   ref=ref, snapshot=snapshot, **kwargs))

    def geo_target(value: Any, ref: str) -> tuple[str | None, str]:
        """Resolve a geography cell to a country/region object id; returns (id, kind)."""
        text = clean(value)
        if not text:
            return None, "none"
        if text.lower() in UN_REGIONS:
            rid = C.make_id("region", C.slugify(text))
            obj("region", rid, text, ref=ref, attrs={"source": "porb"},
                alt_ids=[{"scheme": "porb_region", "value": text}])
            return rid, "region"
        iso2 = C.country_iso2(text)
        if iso2:
            cid = C.make_id("country", iso2.lower())
            obj("country", cid, C.country_name(iso2), ref=ref,
                attrs={"iso2": iso2}, alt_ids=[{"scheme": "iso2", "value": iso2}])
            return cid, "country"
        return None, "unresolved"

    def institution(name: str, *, ref: str, extra: dict[str, Any] | None = None) -> str:
        hit = matcher.match(name)
        if hit and hit.get("clarisa_id"):
            iid = C.make_id("institution", f"clarisa-{hit['clarisa_id']}")
            attrs = {"acronym": hit.get("acronym"), "matched": True,
                     "matched_via": "clarisa name/acronym", "porb_name": name}
            alt = [{"scheme": "clarisa_institution_id", "value": str(hit["clarisa_id"])}]
            label = hit.get("name") or name
        else:
            iid = C.make_id("institution", f"porb-{C.slugify(name)}")
            attrs = {"matched": False, "porb_name": name}
            alt = [{"scheme": "porb_partner_name", "value": name}]
            label = name
        obj("institution", iid, label, ref=ref, attrs={**attrs, **(extra or {})}, alt_ids=alt)
        return iid

    # --- programs ---------------------------------------------------------------------
    program_ids: dict[str, str] = {}
    for p in sheets.get("_programs", []):
        code = PROGRAM_CODES.get(p["name"].lower().strip())
        if not code:
            continue
        pid = C.make_id("program", code.lower())
        program_ids[p["name"].lower().strip()] = pid
        obj("program", pid, p["name"], ref=f"PORB master: Program/Accelerator={p['name']}",
            description=f"CGIAR Science Program / Accelerator {code} — {p['name']}.",
            attrs={"code": code, "porb_name": p["name"]},
            alt_ids=[{"scheme": "clarisa_initiative_code", "value": code}])
    prog_id = C.make_id("program", program_code.lower())

    # --- HLO sheet: aow, hlo, indicator ----------------------------------------------
    aows: dict[str, dict[str, Any]] = {}
    hlos: dict[str, dict[str, Any]] = {}
    indicators: list[dict[str, Any]] = []
    hlo_budget: dict[str, float] = defaultdict(float)
    hlo_centres: dict[str, set[str]] = defaultdict(set)
    hlo_rows: dict[str, int] = defaultdict(int)
    skipped_subtotals = 0

    for i, row in enumerate(sheets.get("HLO", []), start=2):
        hlo_raw = row.get("High Level Output")
        if is_subtotal(hlo_raw):
            skipped_subtotals += 1
            continue
        aow = parse_aow(row.get("AOW"), program_code)
        hlo = parse_hlo(hlo_raw, program_code, aow["code"] if aow else None)
        if not hlo:
            continue
        if aow:
            aows.setdefault(aow["id"], aow)
        hlos.setdefault(hlo["id"], hlo)
        centre = clean(row.get("Center"))
        budget = money(row.get("Budget (USD)"))
        kpi = clean(row.get("KPI Type"))
        countries = split_list(row.get("Country(ies) of implementation"))
        key = C.sha1_12("|".join([program_code, str(aow and aow["code"]), str(centre),
                                  str(hlo["code"] or hlo["label"]), str(kpi),
                                  C.trim(row.get("Description"), 300) or "", str(i)]))
        ind = {
            "id": C.make_id("indicator", f"porb-{key}"),
            "hlo_id": hlo["id"], "aow_id": aow["id"] if aow else None,
            "center": centre, "kpi_type": kpi,
            "label": f"{centre or 'Unknown Center'} — {hlo['code'] or hlo['label']} ({kpi or 'custom'})",
            "description": C.trim(row.get("Description"), 2000),
            "target": row.get("Target"), "budget_usd": budget,
            "assumption": C.trim(row.get("Assumption"), 500),
            "countries": countries, "row": i,
        }
        indicators.append(ind)
        if budget:
            hlo_budget[hlo["id"]] += budget
        if centre:
            hlo_centres[hlo["id"]].add(centre)
        hlo_rows[hlo["id"]] += 1

    for aow in aows.values():
        obj("aow", aow["id"], aow["label"], ref=f"PORB HLO sheet: AOW={aow['raw']}",
            attrs={"code": aow["code"], "program_code": program_code},
            alt_ids=[{"scheme": "porb_aow_code", "value": aow["code"]}])
        link(aow["id"], "PART_OF", prog_id, f"PORB HLO sheet: AOW={aow['raw']}", band="high")

    for hlo in hlos.values():
        aow_id = next((a["id"] for a in aows.values() if a["code"] == hlo["aow_code"]), None)
        obj("hlo", hlo["id"], hlo["label"], ref=f"PORB HLO sheet: {hlo['raw']}",
            description=hlo["raw"],
            attrs={"code": hlo["code"], "aow_code": hlo["aow_code"],
                   "program_code": program_code,
                   # money boundary: budget lives HERE, never on a link (CONTRACT §3.6)
                   "budget_usd": round(hlo_budget.get(hlo["id"], 0.0), 2),
                   "budget_basis": "sum of the PORB HLO-sheet rows of this HLO",
                   "n_indicator_rows": hlo_rows.get(hlo["id"], 0),
                   "centres": sorted(hlo_centres.get(hlo["id"], []))},
            alt_ids=[{"scheme": "porb_hlo_code", "value": hlo["code"]}] if hlo["code"] else None)
        if aow_id:
            link(hlo["id"], "PART_OF", aow_id, f"PORB HLO sheet: {hlo['raw']}", band="high")

    for ind in indicators:
        obj("indicator", ind["id"], ind["label"], ref=f"PORB HLO sheet row {ind['row']}",
            description=ind["description"],
            attrs={"center": ind["center"], "kpi_type": ind["kpi_type"],
                   "target": ind["target"], "budget_usd": ind["budget_usd"],
                   "assumption": ind["assumption"], "program_code": program_code,
                   "countries_of_implementation": ind["countries"]},
            alt_ids=[{"scheme": "porb_row", "value": f"HLO!{ind['row']}"}])
        link(ind["id"], "PART_OF", ind["hlo_id"], f"PORB HLO sheet row {ind['row']}", band="high")
        for cname in ind["countries"]:
            tid, kind = geo_target(cname, f"PORB HLO sheet row {ind['row']}")
            if tid:
                link(ind["id"], "LOCATED_IN", tid, f"PORB HLO sheet row {ind['row']}",
                     attrs={"kind": kind, "source_sheet": "HLO"})

    # --- Outcomes sheet ---------------------------------------------------------------
    outcomes: dict[str, dict[str, Any]] = {}
    for i, row in enumerate(sheets.get("Outcomes", []), start=2):
        oc = parse_outcome(row.get("Outcome"), program_code)
        if not oc:
            continue
        rec = outcomes.setdefault(oc["id"], {**oc, "indicator_types": set(), "targets": [],
                                             "geos": set(), "types": set(), "aows": set()})
        if clean(row.get("Indicator Type")):
            rec["indicator_types"].add(clean(row.get("Indicator Type")))
        if clean(row.get("Type of Outcome")):
            rec["types"].add(clean(row.get("Type of Outcome")))
        if clean(row.get("Geographic Location")):
            rec["geos"].add(clean(row.get("Geographic Location")))
        if row.get("Target Value") is not None:
            rec["targets"].append(row.get("Target Value"))
        aow = parse_aow(row.get("AOW"), program_code)
        if aow:
            rec["aows"].add(aow["id"])

    for oc in outcomes.values():
        horizon = "2030" if (oc["code"] or "").startswith("2030") else "interim"
        obj("outcome", oc["id"], oc["label"], ref=f"PORB Outcomes sheet: {oc['code'] or oc['label'][:60]}",
            description=oc["raw"],
            attrs={"code": oc["code"], "horizon": horizon,
                   "outcome_types": sorted(oc["types"]),
                   "indicator_types": sorted(oc["indicator_types"]),
                   "geographies": sorted(oc["geos"]),
                   "targets": [t for t in oc["targets"] if t is not None][:10],
                   "program_code": program_code},
            alt_ids=[{"scheme": "porb_outcome_code", "value": oc["code"]}] if oc["code"] else None)
        link(oc["id"], "PART_OF", prog_id, "PORB Outcomes sheet", band="high")
        for geo in oc["geos"]:
            for piece in re.split(r"[:,]", geo.replace("Country", "")):
                tid, kind = geo_target(piece, "PORB Outcomes sheet")
                if tid:
                    link(oc["id"], "LOCATED_IN", tid, "PORB Outcomes sheet",
                         attrs={"kind": kind, "source_sheet": "Outcomes"})

    def resolve_framework_target(text: Any) -> tuple[str | None, str]:
        """Resolve a free-text 'High Level Output title' / 'Supported outcomes' cell."""
        t = clean(text)
        if not t:
            return None, "none"
        oc = parse_outcome(t, program_code)
        if oc and oc["code"] and oc["id"] in outcomes:
            return oc["id"], "outcome_code"
        hl = parse_hlo(t, program_code, None)
        if hl and hl["id"] in hlos:
            return hl["id"], "hlo_code"
        norm = C.slugify(t)[:40]
        for hid, h in hlos.items():
            if norm and norm in C.slugify(h["raw"]):
                return hid, "hlo_title"
        for oid, o in outcomes.items():
            if norm and norm in C.slugify(o["raw"]):
                return oid, "outcome_title"
        return None, "unmatched"

    # --- W3-Bilateral sheet: projects -------------------------------------------------
    projects: list[dict[str, Any]] = []
    unmatched_funding = 0
    for i, row in enumerate(sheets.get("W3-Bilateral", []), start=2):
        title = clean(row.get("Project title"))
        if not title or is_subtotal(title, row.get("High Level Output title")):
            skipped_subtotals += 1 if title else 0
            continue
        centre = clean(row.get("Center"))
        amount = money(row.get("W3/Bilateral Project (USD)"))
        pid = C.make_id("project", "porb-" + C.sha1_12(f"{program}|{centre}|{title}"))
        projects.append({"id": pid, "title": title, "center": centre, "amount_usd": amount,
                         "row": i, "framework": clean(row.get("High Level Output title"))})
        obj("project", pid, title, ref=f"PORB W3-Bilateral row {i}",
            attrs={"center": centre, "amount_usd": amount, "funding_type": "W3/Bilateral",
                   "program_code": program_code,
                   "assumption": C.trim(row.get("Assumption"), 300)},
            alt_ids=[{"scheme": "porb_row", "value": f"W3-Bilateral!{i}"}])
        target, how = resolve_framework_target(row.get("High Level Output title"))
        if target:
            link(target, "FUNDED_BY", pid, f"PORB W3-Bilateral row {i}",
                 attrs={"amount_usd": amount, "matched_on": how,
                        "reported_against": clean(row.get("High Level Output title"))},
                 band="high" if how.endswith("code") else "medium")
        else:
            unmatched_funding += 1
            link(prog_id, "FUNDED_BY", pid, f"PORB W3-Bilateral row {i}",
                 attrs={"amount_usd": amount, "matched_on": "program_fallback",
                        "reported_against": clean(row.get("High Level Output title"))},
                 band="medium",
                 note="The reported framework target could not be resolved to an HLO/outcome.")
        if centre:
            institution(centre, ref=f"PORB W3-Bilateral row {i}", extra={"role": "managing centre"})

    # --- MELIA sheet ------------------------------------------------------------------
    melia: list[dict[str, Any]] = []
    for i, row in enumerate(sheets.get("MELIA", []), start=2):
        title = clean(row.get("MELIA study"))
        if not title or is_subtotal(title):
            skipped_subtotals += 1 if title else 0
            continue
        mid = C.make_id("melia_study", "porb-" + C.sha1_12(f"{program}|{clean(row.get('Center'))}|{title}"))
        melia.append({"id": mid, "title": title, "row": i})
        obj("melia_study", mid, title, ref=f"PORB MELIA row {i}",
            attrs={"center": clean(row.get("Center")), "aow": clean(row.get("AOW")),
                   "budget_usd": money(row.get("Total Budget (USD)")),
                   "geographic_location": clean(row.get("Geographic location")),
                   "program_code": program_code},
            alt_ids=[{"scheme": "porb_row", "value": f"MELIA!{i}"}])
        target, how = resolve_framework_target(row.get("Supported outcomes"))
        if target:
            link(target, "STUDIED_BY", mid, f"PORB MELIA row {i}",
                 attrs={"matched_on": how, "reported_against": clean(row.get("Supported outcomes"))},
                 band="high" if how.endswith("code") else "medium")
        tid, kind = geo_target(row.get("Geographic location"), f"PORB MELIA row {i}")
        if tid:
            link(mid, "LOCATED_IN", tid, f"PORB MELIA row {i}", attrs={"kind": kind})

    # --- Partners sheet ---------------------------------------------------------------
    partner_rows = 0
    for i, row in enumerate(sheets.get("Partners", []), start=2):
        name = clean(row.get("Partner"))
        if not name or is_subtotal(name):
            skipped_subtotals += 1 if name else 0
            continue
        partner_rows += 1
        aow = parse_aow(row.get("AOW"), program_code)
        subject = aow["id"] if aow and aow["id"] in seen else prog_id
        iid = institution(name, ref=f"PORB Partners row {i}",
                          extra={"geographic_location": clean(row.get("Geographic location"))})
        link(subject, "WITH_PARTNER", iid, f"PORB Partners row {i}",
             attrs={"budget_usd": money(row.get("Total Budget (USD)")),
                    "geographic_location": clean(row.get("Geographic location")),
                    "centre": clean(row.get("Center")), "granularity": "aow",
                    "assumption": C.trim(row.get("Assumption"), 300)},
             band="high")
        tid, kind = geo_target(row.get("Geographic location"), f"PORB Partners row {i}")
        if tid:
            link(iid, "LOCATED_IN", tid, f"PORB Partners row {i}", attrs={"kind": kind})

    # --- Countries of Implementation / Location of Benefit ----------------------------
    geo_rows = {"implementation": 0, "benefit": 0, "global": 0, "unresolved": 0}
    global_share: dict[str, float] = defaultdict(float)
    for sheet, kind, col in (("Countries of Implementation", "implementation", "Country"),
                             ("Location of Benefit", "benefit", "Location")):
        for i, row in enumerate(sheets.get(sheet, []), start=2):
            aow = parse_aow(row.get("AOW"), program_code)
            if not aow or aow["id"] not in seen:
                continue
            value = clean(row.get(col))
            share = money(row.get("Percentage (%)"))
            tid, gkind = geo_target(value, f"PORB {sheet} row {i}")
            if tid:
                geo_rows[kind] += 1
                link(aow["id"], "LOCATED_IN", tid, f"PORB {sheet} row {i}",
                     attrs={"share_pct": share, "kind": kind, "geo_kind": gkind,
                            "centre": clean(row.get("Center")), "source_sheet": sheet},
                     disambiguator=kind, band="high")
            elif value and value.lower() in C.NON_COUNTRY_GEO:
                geo_rows["global"] += 1
                global_share[aow["id"]] += share or 0.0
            else:
                geo_rows["unresolved"] += 1
    for aid, share in global_share.items():
        claims.append(C.attr_claim(aid, {"global_share_pct": round(share, 2)},
                                   channel=CHANNEL, system=SYSTEM,
                                   ref="PORB Countries of Implementation / Location of Benefit",
                                   snapshot=snapshot,
                                   note="'Global'/'Regional' shares are an attribute, not a link."))

    # --- Synergy Programs -------------------------------------------------------------
    synergy_rows = 0
    for i, row in enumerate(sheets.get("Synergy Programs", []), start=2):
        other = clean(row.get("Program or Accelerator"))
        code = PROGRAM_CODES.get((other or "").lower())
        if not code:
            continue
        target_prog = C.make_id("program", code.lower())
        hl = parse_hlo(row.get("High Level Output"), program_code, None)
        subject = hl["id"] if hl and hl["id"] in seen else prog_id
        synergy_rows += 1
        link(subject, "SYNERGY_WITH", target_prog, f"PORB Synergy Programs row {i}",
             attrs={"description": C.trim(row.get("Brief description"), 600),
                    "aow": clean(row.get("AOW"))}, band="high")

    claims = C.dedupe_links(C.dedupe_objects(claims))

    C.write_extract(extract, {
        "channel": CHANNEL, "snapshot": snapshot, "source_file": xlsx, "public": True,
        "program": {"name": program, "code": program_code}, "generated_at": C.now_iso(),
        "counts": {"aows": len(aows), "hlos": len(hlos), "indicators": len(indicators),
                   "outcomes": len(outcomes), "projects": len(projects), "melia": len(melia),
                   "partner_rows": partner_rows, "synergy_rows": synergy_rows,
                   "skipped_subtotal_rows": skipped_subtotals,
                   "unmatched_funding_rows": unmatched_funding, "geo_rows": geo_rows},
        "aows": list(aows.values()),
        "hlos": [{**h, "budget_usd": round(hlo_budget.get(h["id"], 0.0), 2),
                  "n_indicator_rows": hlo_rows.get(h["id"], 0)} for h in hlos.values()],
        "indicators": indicators,
        "outcomes": [{k: (sorted(v) if isinstance(v, set) else v) for k, v in o.items()}
                     for o in outcomes.values()],
        "projects": projects,
        "melia_studies": melia,
    })
    n = C.write_batch(out, claims)
    summary = C.summarise(claims)
    summary.update(channel=CHANNEL, out=out, extract=extract, claims=n)
    return summary


def run_channel(limit: int | None = None, load: bool = True, out: str | None = None,
                program: str = "Breeding for Tomorrow", **_: object) -> dict:
    """Keyword-only entry point for POST /api/ingest/porb (PLAN §4). `limit` is not meaningful
    for a workbook channel (the whole programme sheet set is one consistent unit) and is ignored."""
    out_path = C.channel_out_path(CHANNEL, out)
    return C.channel_entry(CHANNEL, lambda: run(
        DEFAULT_XLSX, program, out_path, C.runtime_extract_path(CHANNEL),
        "data/seed/prms_sp01_results.json"), load=load)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="TRACE intake channel: PORB")
    ap.add_argument("--xlsx", default=DEFAULT_XLSX)
    ap.add_argument("--program", default="Breeding for Tomorrow")
    ap.add_argument("--out", default="data/claims/seed-porb.jsonl")
    ap.add_argument("--extract", default="data/seed/porb_sp01.json")
    ap.add_argument("--prms-extract", default="data/seed/prms_sp01_results.json")
    a = ap.parse_args(argv)
    s = run(a.xlsx, a.program, a.out, a.extract, a.prms_extract)
    print(f"[porb] {s['claims']} claims -> {s['out']}")
    print(f"       objects: {s['objects']}")
    print(f"       links:   {s['links']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
